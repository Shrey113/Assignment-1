import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

PPTX_PATH = "Smart_CPU_Scheduler_Presentation.pptx"

# ==============================================================================
# DESIGN SYSTEM: MINIMALIST, SPACIOUS, MODERN (NO HEAVY BOXES)
# ==============================================================================
COLOR_BG_WHITE = RGBColor(255, 255, 255)       # Clean White (#FFFFFF)
COLOR_PRIMARY_NAVY = RGBColor(15, 23, 42)      # Slate 900 (#0F172A)
COLOR_TEXT_MAIN = RGBColor(30, 41, 59)         # Slate 800 (#1E293B)
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)     # Slate 500 (#64748B)
COLOR_TEXT_LIGHT = RGBColor(148, 163, 184)     # Slate 400 (#94A3B8)
COLOR_ACCENT_BLUE = RGBColor(37, 99, 235)      # Blue 600 (#2563EB)
COLOR_ACCENT_EMERALD = RGBColor(16, 185, 129)  # Emerald 500 (#10B981)
COLOR_DIVIDER = RGBColor(226, 232, 240)        # Slate 200 (#E2E8F0)
COLOR_IMG_BORDER = RGBColor(203, 213, 225)     # Slate 300 (#CBD5E1) - subtle crisp frame

FONT_NAME = "Calibri"

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG_WHITE

def add_header(slide, slide_num, category, title, subtitle):
    # Category / Super-title
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(11.7), Inches(0.25))
    ctf = cat_box.text_frame
    ctf.word_wrap = True
    ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = 0
    p_cat = ctf.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.name = FONT_NAME
    p_cat.font.size = Pt(9.5)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_ACCENT_BLUE

    # Title & Subtitle in single clean block
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.7), Inches(0.85))
    ttf = title_box.text_frame
    ttf.word_wrap = True
    ttf.margin_left = ttf.margin_top = ttf.margin_right = ttf.margin_bottom = 0
    
    p_title = ttf.paragraphs[0]
    p_title.text = title
    p_title.font.name = FONT_NAME
    p_title.font.size = Pt(21)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_PRIMARY_NAVY
    p_title.space_after = Pt(2)

    p_sub = ttf.add_paragraph()
    p_sub.text = subtitle
    p_sub.font.name = FONT_NAME
    p_sub.font.size = Pt(11)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED

    # Subtle horizontal divider line (clean separator, no box)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.58), Inches(11.733), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_DIVIDER
    line.line.fill.background()

    # Footer Left
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(9.5), Inches(0.3))
    ftf = footer_box.text_frame
    ftf.word_wrap = True
    ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0
    p_foot = ftf.paragraphs[0]
    p_foot.text = "Smart CPU Scheduler Simulator  •  Operating Systems Innovative Assignment 1  •  Nirma University"
    p_foot.font.name = FONT_NAME
    p_foot.font.size = Pt(9)
    p_foot.font.color.rgb = COLOR_TEXT_MUTED

    # Footer Right (Slide number)
    num_box = slide.shapes.add_textbox(Inches(10.5), Inches(7.05), Inches(2.0), Inches(0.3))
    ntf = num_box.text_frame
    ntf.word_wrap = False
    ntf.margin_left = ntf.margin_top = ntf.margin_right = ntf.margin_bottom = 0
    p_num = ntf.paragraphs[0]
    p_num.text = f"{slide_num} / 5"
    p_num.alignment = PP_ALIGN.RIGHT
    p_num.font.name = FONT_NAME
    p_num.font.size = Pt(9.5)
    p_num.font.bold = True
    p_num.font.color.rgb = COLOR_ACCENT_BLUE

def add_clean_screenshot(slide, img_path, left, top, width, height, caption=""):
    """Places screenshot directly on the slide with a clean subtle border (no heavy box container)"""
    if os.path.exists(img_path):
        pic = slide.shapes.add_picture(img_path, left, top, width, height)
        # Apply a subtle, clean single-line border
        pic.line.color.rgb = COLOR_IMG_BORDER
        pic.line.width = Pt(0.75)
    
    if caption:
        cap_box = slide.shapes.add_textbox(left, top + height + Inches(0.08), width, Inches(0.32))
        ctf = cap_box.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = 0
        p = ctf.paragraphs[0]
        p.text = caption
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_NAME
        p.font.size = Pt(8.5)
        p.font.italic = True
        p.font.color.rgb = COLOR_TEXT_MUTED

# ==============================================================================
# SLIDE 1: CLEAN COVER & ACADEMIC DETAILS (NO BOXES)
# ==============================================================================
def build_slide_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    # Nirma Logo (Top Left)
    logo_path = "data/temp/Nirma.png"
    if os.path.exists(logo_path):
        slide.shapes.add_picture(logo_path, Inches(0.8), Inches(0.65), Inches(2.2), Inches(0.93))

    # University Institution Heading (Right of Logo)
    inst_box = slide.shapes.add_textbox(Inches(3.25), Inches(0.72), Inches(9.2), Inches(0.8))
    itf = inst_box.text_frame
    itf.word_wrap = True
    itf.margin_left = itf.margin_top = itf.margin_right = itf.margin_bottom = 0
    
    p1 = itf.paragraphs[0]
    p1.text = "NIRMA UNIVERSITY  •  INSTITUTE OF TECHNOLOGY"
    p1.font.name = FONT_NAME
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_PRIMARY_NAVY
    p1.space_after = Pt(2)

    p2 = itf.add_paragraph()
    p2.text = "Department of Computer Science & Engineering  |  Master of Computer Applications (MCA)"
    p2.font.name = FONT_NAME
    p2.font.size = Pt(10)
    p2.font.color.rgb = COLOR_TEXT_MUTED

    # Top clean divider line
    div1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.85), Inches(11.733), Inches(0.015))
    div1.fill.solid()
    div1.fill.fore_color.rgb = COLOR_DIVIDER
    div1.line.fill.background()

    # Main Presentation Title Block
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.05), Inches(11.733), Inches(1.35))
    ttf = title_box.text_frame
    ttf.word_wrap = True
    ttf.margin_left = ttf.margin_top = ttf.margin_right = ttf.margin_bottom = 0

    p_badge = ttf.paragraphs[0]
    p_badge.text = "OPERATING SYSTEMS INNOVATIVE ASSIGNMENT – 1 (6CS102CC25)"
    p_badge.font.name = FONT_NAME
    p_badge.font.size = Pt(10)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_ACCENT_BLUE
    p_badge.space_after = Pt(4)

    p_main = ttf.add_paragraph()
    p_main.text = "Smart CPU Scheduler Simulator"
    p_main.font.name = FONT_NAME
    p_main.font.size = Pt(30)
    p_main.font.bold = True
    p_main.font.color.rgb = COLOR_PRIMARY_NAVY
    p_main.space_after = Pt(4)

    p_desc = ttf.add_paragraph()
    p_desc.text = "Interactive Gantt Visualization, Multi-Algorithm Benchmark Comparison & Intelligent Workload Recommender"
    p_desc.font.name = FONT_NAME
    p_desc.font.size = Pt(12)
    p_desc.font.color.rgb = COLOR_TEXT_MUTED

    # Middle clean divider line
    div2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.55), Inches(11.733), Inches(0.015))
    div2.fill.solid()
    div2.fill.fore_color.rgb = COLOR_DIVIDER
    div2.line.fill.background()

    # Left Column: Project Highlights (Clean typography, no box)
    left_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(6.0), Inches(3.2))
    ltf = left_box.text_frame
    ltf.word_wrap = True
    ltf.margin_left = ltf.margin_top = ltf.margin_right = ltf.margin_bottom = 0

    lp_head = ltf.paragraphs[0]
    lp_head.text = "Project Objectives & Key Features"
    lp_head.font.name = FONT_NAME
    lp_head.font.size = Pt(14)
    lp_head.font.bold = True
    lp_head.font.color.rgb = COLOR_PRIMARY_NAVY
    lp_head.space_after = Pt(8)

    highlights = [
        ("Complete Algorithmic Suite: ", "Implements FCFS, SJF (Non-Preemptive), SRTF (Preemptive), Priority (Preemptive & Non-Preemptive), and Round Robin."),
        ("Interactive Web Dashboard: ", "Built with Streamlit & Plotly for dynamic process inputs, real-time metrics, and color-coded Gantt timelines."),
        ("Multi-Algorithm Comparison: ", "Side-by-side benchmark bar charts to visually and empirically analyze the Convoy Effect and scheduling trade-offs."),
        ("Smart Heuristic Recommender: ", "Statistical workload analysis (burst variance & priority spread) to dynamically recommend the optimal scheduling algorithm.")
    ]
    for b_title, b_desc in highlights:
        p = ltf.add_paragraph()
        p.text = "• "
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_ACCENT_BLUE
        p.space_after = Pt(6)
        r1 = p.add_run()
        r1.text = b_title
        r1.font.bold = True
        r1.font.color.rgb = COLOR_TEXT_MAIN
        r2 = p.add_run()
        r2.text = b_desc
        r2.font.bold = False
        r2.font.color.rgb = COLOR_TEXT_MUTED

    # Right Column: Academic & Submission Details (Clean alignment, no box)
    right_box = slide.shapes.add_textbox(Inches(7.2), Inches(3.8), Inches(5.3), Inches(3.2))
    rtf = right_box.text_frame
    rtf.word_wrap = True
    rtf.margin_left = rtf.margin_top = rtf.margin_right = rtf.margin_bottom = 0

    rp_head = rtf.paragraphs[0]
    rp_head.text = "Submission & Author Credentials"
    rp_head.font.name = FONT_NAME
    rp_head.font.size = Pt(14)
    rp_head.font.bold = True
    rp_head.font.color.rgb = COLOR_PRIMARY_NAVY
    rp_head.space_after = Pt(8)

    meta_rows = [
        ("Student Name", "Shrey Patel"),
        ("Roll Number", "26MCA042"),
        ("Program & Semester", "Master of Computer Applications (MCA - Sem 1)"),
        ("Subject Name", "Operating System (OS)"),
        ("Subject Code", "6CS102CC25"),
        ("Faculty Guide", "Drashti Patel"),
        ("Academic Year", "2026 – 2027"),
        ("Technology Stack", "Python 3.11, Streamlit, Plotly, Pandas")
    ]
    for label, val in meta_rows:
        p = rtf.add_paragraph()
        p.space_after = Pt(4)
        run_l = p.add_run()
        run_l.text = f"{label:20}:  "
        run_l.font.name = "Consolas"
        run_l.font.size = Pt(9.5)
        run_l.font.bold = True
        run_l.font.color.rgb = COLOR_ACCENT_BLUE

        run_v = p.add_run()
        run_v.text = val
        run_v.font.name = FONT_NAME
        run_v.font.size = Pt(10)
        run_v.font.bold = (label in ["Student Name", "Roll Number"])
        run_v.font.color.rgb = COLOR_PRIMARY_NAVY if label in ["Student Name", "Roll Number"] else COLOR_TEXT_MAIN

# ==============================================================================
# SLIDE 2: SINGLE ALGORITHM SIMULATION & GANTT VISUALIZER (NO BOXES)
# ==============================================================================
def build_slide_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    add_header(slide, 2, "MODULE 1: INTERACTIVE SIMULATION & TIMELINE", 
               "Single Algorithm Simulation & Gantt Visualizer", 
               "Automated timeline scheduling, real-time KPI metrics, and responsive Plotly Gantt charts")

    # Left Column: Clean bullet text with generous whitespace (width: 4.3 in)
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.3), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    # Section 1: Capabilities
    p1 = tf.paragraphs[0]
    p1.text = "Interactive Capabilities"
    p1.font.name = FONT_NAME
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_PRIMARY_NAVY
    p1.space_after = Pt(6)

    bullets1 = [
        ("Configurable Inputs: ", "Add custom processes or select pre-configured workloads (Short Jobs, Mixed Jobs, Priority Spreads)."),
        ("Dynamic Parameters: ", "Fine-tune Round Robin Quantum Q, Preemptive vs Non-Preemptive toggles in real-time."),
        ("Plotly Gantt Chart: ", "Color-coded visual bars showing exact execution time slices, context switches, and CPU idle intervals.")
    ]
    for b_title, b_desc in bullets1:
        p = tf.add_paragraph()
        p.text = "• "
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_ACCENT_BLUE
        p.space_after = Pt(5)
        r1 = p.add_run()
        r1.text = b_title
        r1.font.bold = True
        r1.font.color.rgb = COLOR_TEXT_MAIN
        r2 = p.add_run()
        r2.text = b_desc
        r2.font.bold = False
        r2.font.color.rgb = COLOR_TEXT_MUTED

    # Spacing between sections
    p_space = tf.add_paragraph()
    p_space.space_after = Pt(10)

    # Section 2: Core Metrics
    p2 = tf.add_paragraph()
    p2.text = "Performance Metrics Computed"
    p2.font.name = FONT_NAME
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_PRIMARY_NAVY
    p2.space_after = Pt(6)

    kpis = [
        ("Turnaround Time (TAT): ", "Finish Time − Arrival Time (Total system duration)."),
        ("Waiting Time (WT): ", "Turnaround Time − Burst Time (Delay in Ready Queue)."),
        ("Response Time (RT): ", "First CPU Execution − Arrival Time (Latency)."),
        ("CPU Utilization (%): ", "(Busy CPU Time / Total Schedule Length) * 100."),
        ("Context Switches: ", "Total process preemptions & state switches logged.")
    ]
    for k_title, k_desc in kpis:
        p = tf.add_paragraph()
        p.text = "✔ "
        p.font.size = Pt(9)
        p.font.color.rgb = COLOR_ACCENT_EMERALD
        p.space_after = Pt(4)
        r1 = p.add_run()
        r1.text = k_title
        r1.font.bold = True
        r1.font.color.rgb = COLOR_TEXT_MAIN
        r2 = p.add_run()
        r2.text = k_desc
        r2.font.bold = False
        r2.font.color.rgb = COLOR_TEXT_MUTED

    # Right Column: Clean, Large Screenshot (No enclosing rounded box!)
    # Width: 7.3 in, Height: 7.3 / 1.5238 = 4.79 in
    img_path = "screenshots/01_single_algorithm_gantt.png"
    add_clean_screenshot(slide, img_path, Inches(5.3), Inches(1.8), Inches(7.233), Inches(4.75),
                         caption="Figure 1: Live Web App — Process Table, KPI Metric Cards, and Plotly Gantt Execution Timeline")

# ==============================================================================
# SLIDE 3: MULTI-ALGORITHM COMPARATIVE BENCHMARK (NO BOXES)
# ==============================================================================
def build_slide_3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    add_header(slide, 3, "MODULE 2: BENCHMARKING & CONVOY EFFECT ANALYSIS", 
               "Multi-Algorithm Comparative Benchmark", 
               "Side-by-side empirical performance comparison across all 5 algorithms on identical workloads")

    # Left Column: Insights & Analysis (width: 4.3 in)
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.3), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    # Section 1: Simultaneous Benchmark
    p1 = tf.paragraphs[0]
    p1.text = "Simultaneous Multi-Algorithm Evaluation"
    p1.font.name = FONT_NAME
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_PRIMARY_NAVY
    p1.space_after = Pt(6)

    comp_points = [
        ("Side-by-Side Execution: ", "Runs FCFS, SJF, SRTF, Priority, and Round Robin simultaneously under identical process arrivals."),
        ("Grouped Bar Charts: ", "Plots Avg Waiting Time, Turnaround Time, and Response Time together for instant visual contrast."),
        ("Empirical Validation: ", "Provides quantitative proof of theoretical scheduling trade-offs directly in the browser.")
    ]
    for b_title, b_desc in comp_points:
        p = tf.add_paragraph()
        p.text = "• "
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_ACCENT_BLUE
        p.space_after = Pt(5)
        r1 = p.add_run()
        r1.text = b_title
        r1.font.bold = True
        r1.font.color.rgb = COLOR_TEXT_MAIN
        r2 = p.add_run()
        r2.text = b_desc
        r2.font.bold = False
        r2.font.color.rgb = COLOR_TEXT_MUTED

    # Spacing between sections
    p_space = tf.add_paragraph()
    p_space.space_after = Pt(10)

    # Section 2: Convoy Effect Analysis
    p2 = tf.add_paragraph()
    p2.text = "Convoy Effect Case Study (Test Case 2)"
    p2.font.name = FONT_NAME
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_ACCENT_BLUE
    p2.space_after = Pt(6)

    findings = [
        ("FCFS Convoy Penalty: ", "P1 (Burst: 8ms) blocks ready short jobs, escalating Avg Waiting Time to 8.80 ms."),
        ("SRTF Preemption Win: ", "Preempts long jobs for shorter tasks, slashing Avg WT to 3.60 ms (59% reduction!)."),
        ("Round Robin Balance: ", "Maintains balanced response latency (2.40 ms) across all processes, preventing starvation.")
    ]
    for f_title, f_desc in findings:
        p = tf.add_paragraph()
        p.text = "⚡ "
        p.font.size = Pt(9)
        p.space_after = Pt(4)
        r1 = p.add_run()
        r1.text = f_title
        r1.font.bold = True
        r1.font.color.rgb = COLOR_TEXT_MAIN
        r2 = p.add_run()
        r2.text = f_desc
        r2.font.bold = False
        r2.font.color.rgb = COLOR_TEXT_MUTED

    # Right Column: Large, Expansive Screenshot of the Website Comparison (No boxes!)
    img_path = "screenshots/02_comparative_benchmark.png"
    add_clean_screenshot(slide, img_path, Inches(5.3), Inches(1.8), Inches(7.233), Inches(4.75),
                         caption="Figure 2: Live Web App — Multi-Algorithm Benchmark Comparison (Grouped Bar Charts & Performance Table)")

# ==============================================================================
# SLIDE 4: INTELLIGENT WORKLOAD RECOMMENDER INNOVATION (NO BOXES)
# ==============================================================================
def build_slide_4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    add_header(slide, 4, "MODULE 3: ADDITIONAL INNOVATION FEATURE", 
               "Intelligent Workload-Based Algorithm Recommender", 
               "Heuristic AI engine that dynamically analyzes batch characteristics to select the optimal scheduling policy")

    # Left Column: Decision Logic (width: 4.3 in)
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.3), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p1 = tf.paragraphs[0]
    p1.text = "Heuristic Decision Rules"
    p1.font.name = FONT_NAME
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_PRIMARY_NAVY
    p1.space_after = Pt(6)

    heuristics = [
        ("High Burst Variance (> 4.0): ", "Detects severe burst heterogeneity. Recommends Round Robin or SRTF to eliminate FCFS convoy delays."),
        ("Homogeneous Short Tasks (≤ 4ms): ", "Detects quick micro-bursts. Recommends SJF / SRTF to mathematically minimize average turnaround time."),
        ("High Priority Spread (≥ 3 Tiers): ", "Detects mission-critical tasks mixed with background jobs. Recommends Preemptive Priority."),
        ("Synchronous Batch Arrivals: ", "Recommends FCFS for uniform sequential pipelines to avoid redundant context switching overhead.")
    ]
    for h_title, h_desc in heuristics:
        p = tf.add_paragraph()
        p.text = "• "
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_ACCENT_BLUE
        p.space_after = Pt(5)
        r1 = p.add_run()
        r1.text = h_title
        r1.font.bold = True
        r1.font.color.rgb = COLOR_TEXT_MAIN
        r2 = p.add_run()
        r2.text = h_desc
        r2.font.bold = False
        r2.font.color.rgb = COLOR_TEXT_MUTED

    p_space = tf.add_paragraph()
    p_space.space_after = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "Explainable Intelligence"
    p2.font.name = FONT_NAME
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_ACCENT_EMERALD
    p2.space_after = Pt(6)

    exp_points = [
        ("Clear Rationale: ", "Explains WHY an algorithm is chosen with confidence metrics."),
        ("One-Click Run: ", "Instantly runs the recommended policy and generates its Gantt chart."),
        ("Pedagogical Value: ", "Helps students connect workload characteristics directly to algorithm behavior.")
    ]
    for e_title, e_desc in exp_points:
        p = tf.add_paragraph()
        p.text = "✔ "
        p.font.size = Pt(9)
        p.font.color.rgb = COLOR_ACCENT_EMERALD
        p.space_after = Pt(3)
        r1 = p.add_run()
        r1.text = e_title
        r1.font.bold = True
        r1.font.color.rgb = COLOR_TEXT_MAIN
        r2 = p.add_run()
        r2.text = e_desc
        r2.font.bold = False
        r2.font.color.rgb = COLOR_TEXT_MUTED

    # Right Column: Large Screenshot of Smart Recommender (No boxes!)
    img_path = "screenshots/03_smart_recommendation.png"
    add_clean_screenshot(slide, img_path, Inches(5.3), Inches(1.8), Inches(7.233), Inches(4.75),
                         caption="Figure 3: Live Web App — Smart Workload Recommender (Workload Insights, Algorithm Ranking & Gantt Verification)")

# ==============================================================================
# SLIDE 5: ARCHITECTURE, RESULTS & CONCLUSION (NO BOXES)
# ==============================================================================
def build_slide_5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    add_header(slide, 5, "SYNTHESIS: ARCHITECTURE, BENCHMARK MATRIX & TAKEAWAYS", 
               "System Architecture, Benchmark Matrix & Key Takeaways", 
               "Summary of decoupled modular code, quantitative test workload results, and core conclusions")

    col_w = Inches(3.75)
    gap = Inches(0.24)
    c1_left = Inches(0.8)
    c2_left = c1_left + col_w + gap
    c3_left = c2_left + col_w + gap
    top_pos = Inches(1.8)

    # -------------------------------------------------------------
    # Column 1: Clean Architecture (No box!)
    # -------------------------------------------------------------
    tb1 = slide.shapes.add_textbox(c1_left, top_pos, col_w, Inches(5.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

    p1 = tf1.paragraphs[0]
    p1.text = "1. Decoupled Architecture"
    p1.font.name = FONT_NAME
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_PRIMARY_NAVY
    p1.space_after = Pt(8)

    modules = [
        ("scheduler/process.py", "Lightweight process structure (PID, Arrival, Burst, Priority, Remaining)."),
        ("scheduler/*.py", "Transparent algorithm engines (FCFS, SJF, SRTF, Priority, RR) with zero class bloat."),
        ("scheduler/metrics.py", "Accurate mathematical computation of WT, TAT, RT, and context switches."),
        ("scheduler/smart_suggest.py", "Workload pattern analyzer and heuristic recommendation engine."),
        ("ui/gantt.py", "Plotly visualizer rendering interactive execution bars and timeline ranges."),
        ("app.py", "Streamlit UI orchestrator with multi-tab navigation and sidebar controls.")
    ]
    for mod_name, mod_desc in modules:
        p = tf1.add_paragraph()
        p.space_after = Pt(5)
        run_m = p.add_run()
        run_m.text = f"• {mod_name}\n"
        run_m.font.name = "Consolas"
        run_m.font.size = Pt(9)
        run_m.font.bold = True
        run_m.font.color.rgb = COLOR_ACCENT_BLUE

        run_d = p.add_run()
        run_d.text = f"  {mod_desc}"
        run_d.font.name = FONT_NAME
        run_d.font.size = Pt(8.5)
        run_d.font.color.rgb = COLOR_TEXT_MUTED

    # -------------------------------------------------------------
    # Column 2: Benchmark Matrix Table & Quantitative Notes (No outer box!)
    # -------------------------------------------------------------
    tb2 = slide.shapes.add_textbox(c2_left, top_pos, col_w, Inches(1.0))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0

    p2 = tf2.paragraphs[0]
    p2.text = "2. Empirical Benchmark Matrix"
    p2.font.name = FONT_NAME
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_PRIMARY_NAVY
    p2.space_after = Pt(2)

    p2_sub = tf2.add_paragraph()
    p2_sub.text = "Test Case 2 (Mixed Convoy Workload: 5 Processes):"
    p2_sub.font.name = FONT_NAME
    p2_sub.font.size = Pt(9.5)
    p2_sub.font.color.rgb = COLOR_TEXT_MUTED

    # Clean Table
    table_shape = slide.shapes.add_table(6, 4, c2_left, top_pos + Inches(0.85), col_w, Inches(2.2))
    table = table_shape.table
    table.columns[0].width = Inches(1.4)
    table.columns[1].width = Inches(0.78)
    table.columns[2].width = Inches(0.78)
    table.columns[3].width = Inches(0.78)

    headers = ["Algorithm", "Avg WT", "Avg TAT", "Avg RT"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY_NAVY
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.name = FONT_NAME
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT

    table_data = [
        ("FCFS", "8.80 ms", "13.20 ms", "8.80 ms"),
        ("SJF (Non-Preempt)", "6.00 ms", "10.40 ms", "6.00 ms"),
        ("SRTF (Preemptive)", "3.60 ms", "8.00 ms", "1.60 ms"),
        ("Priority (Preempt)", "4.80 ms", "9.20 ms", "2.60 ms"),
        ("Round Robin (Q=2)", "7.60 ms", "12.00 ms", "2.40 ms"),
    ]

    for i, row in enumerate(table_data):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.fill.solid()
            if i % 2 == 0:
                cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
            else:
                cell.fill.fore_color.rgb = RGBColor(248, 250, 252)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = FONT_NAME
            p.font.size = Pt(8.5)
            if j == 0:
                p.font.bold = True
                p.font.color.rgb = COLOR_TEXT_MAIN
                p.alignment = PP_ALIGN.LEFT
            else:
                p.font.bold = (i == 2 and j == 1)  # Highlight best WT (SRTF)
                p.font.color.rgb = COLOR_ACCENT_BLUE if (i == 2 and j == 1) else COLOR_TEXT_MAIN
                p.alignment = PP_ALIGN.CENTER

    # Additional text note in col 2 below table
    tb2_note = slide.shapes.add_textbox(c2_left, top_pos + Inches(3.3), col_w, Inches(1.6))
    tf2_note = tb2_note.text_frame
    tf2_note.word_wrap = True
    tf2_note.margin_left = tf2_note.margin_top = tf2_note.margin_right = tf2_note.margin_bottom = 0
    
    p_note = tf2_note.paragraphs[0]
    p_note.text = "Quantitative Benchmark Findings:"
    p_note.font.name = FONT_NAME
    p_note.font.size = Pt(10)
    p_note.font.bold = True
    p_note.font.color.rgb = COLOR_PRIMARY_NAVY
    p_note.space_after = Pt(3)

    q_points = [
        "SRTF provides provably lowest WT (3.60ms), reducing latency by 59% over FCFS.",
        "Round Robin ensures uniform responsiveness (RT: 2.40ms) despite higher context switches (11)."
    ]
    for q in q_points:
        p = tf2_note.add_paragraph()
        p.text = f"• {q}"
        p.font.name = FONT_NAME
        p.font.size = Pt(8.5)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_after = Pt(3)

    # -------------------------------------------------------------
    # Column 3: Conclusions & Pedagogical Takeaways (No box!)
    # -------------------------------------------------------------
    tb3 = slide.shapes.add_textbox(c3_left, top_pos, col_w, Inches(5.0))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    tf3.margin_left = tf3.margin_top = tf3.margin_right = tf3.margin_bottom = 0

    p3 = tf3.paragraphs[0]
    p3.text = "3. Conclusions & Takeaways"
    p3.font.name = FONT_NAME
    p3.font.size = Pt(13)
    p3.font.bold = True
    p3.font.color.rgb = COLOR_PRIMARY_NAVY
    p3.space_after = Pt(8)

    conclusions = [
        ("Bridging Theory & Practice: ", "Replaces error-prone manual paper Gantt chart calculations with a verifiable, interactive simulator."),
        ("Empirical Rigor: ", "Visually and quantitatively confirms theoretical phenomena like the FCFS Convoy Effect and RR time slice sensitivity."),
        ("Intelligent Scheduling: ", "Proves that algorithm selection can be automated through feature extraction and variance analysis."),
        ("Clean Engineering: ", "Demonstrates high code maintainability through modular separation between algorithmic logic and presentation layers."),
        ("Future Roadmap: ", "Extensible to Multi-Level Feedback Queues (MLFQ), multi-core CPU scheduling, and dynamic real-time priority aging.")
    ]
    for c_title, c_desc in conclusions:
        p = tf3.add_paragraph()
        p.text = "✔ "
        p.font.size = Pt(9)
        p.font.color.rgb = COLOR_ACCENT_EMERALD
        p.space_after = Pt(4)
        r1 = p.add_run()
        r1.text = c_title
        r1.font.bold = True
        r1.font.color.rgb = COLOR_TEXT_MAIN
        r2 = p.add_run()
        r2.text = c_desc
        r2.font.bold = False
        r2.font.color.rgb = COLOR_TEXT_MUTED

# ==============================================================================
# MAIN GENERATOR
# ==============================================================================
def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    print("Generating Slide 1: Title & Academic Details (Clean open layout)...")
    build_slide_1(prs)

    print("Generating Slide 2: Single Algorithm Gantt Visualizer...")
    build_slide_2(prs)

    print("Generating Slide 3: Multi-Algorithm Comparative Benchmark (Website Comparison)...")
    build_slide_3(prs)

    print("Generating Slide 4: Smart Workload Recommender Innovation...")
    build_slide_4(prs)

    print("Generating Slide 5: System Architecture, Benchmark Matrix & Conclusions...")
    build_slide_5(prs)

    prs.save(PPTX_PATH)
    print(f"Successfully generated PowerPoint presentation at: {os.path.abspath(PPTX_PATH)}")

if __name__ == "__main__":
    main()
